#!/usr/bin/env python3
"""
vehicle_tyre_fitment_from_au_xlsx.py
=====================================
Oct-30 user ask: rebuild `out/rego/vehicle_tyre_fitment.csv` from a
single authoritative source (`AustraliaFormat1_2025.xlsx`), carrying
over 2026 rows from the existing CSV (the xlsx has no 2026 data).

Inputs (defaults; override via CLI flags):
    out/rego/AustraliaFormat1_2025.xlsx
    out/rego/vehicle_tyre_fitment.csv              (existing, for 2026)

Output:
    out/rego/vehicle_tyre_fitment.csv              (overwritten; backed up)
    out/rego/vehicle_tyre_fitment.csv.bak_YYYYMMDD_HHMM

CSV schema (unchanged - app.py and the SPRF already consume this):
    make, model, gen_start_year, size, position

Design decisions (all confirmed by the user Oct-30):
    * Country = 'Australia' only
    * Original (OE) = 1 only  (skips aftermarket / optional fitments)
    * 2026 rows: ALL carried over from the existing CSV verbatim
    * Model-name collapse: scoped per (make, model, year)
          - all trims+versions share the same (size, position) set
                        -> use just `model`
          - all versions within each trim converge
                        -> use `model-trim`
          - versions genuinely differentiate
                        -> use `model-trim-version`
    * Position mapping:  FA=1 RA=1 -> both
                         FA=1 RA=0 -> front
                         FA=0 RA=1 -> rear
                         FA=0 RA=0 -> both  (treat as "both" since the
                                              row still implies a
                                              fitment; preserves the
                                              existing CSV's default)
    * Dedup: by (make, model_resolved, year, size, position)
    * Backup: timestamped suffix so re-runs never clobber a prior backup
    * Load/Speed/Section Width/Aspect Ratio/Rim Size etc. are DROPPED -
      the current 5-column schema is what the SPRF reads.

Usage (from E:\\01. work\\2025\\Data_Anal_Website):
    python vehicle_tyre_fitment_from_au_xlsx.py
    python vehicle_tyre_fitment_from_au_xlsx.py --dry-run
    python vehicle_tyre_fitment_from_au_xlsx.py \
        --xlsx out/rego/AustraliaFormat1_2025.xlsx \
        --existing-csv out/rego/vehicle_tyre_fitment.csv \
        --output out/rego/vehicle_tyre_fitment.csv

Requires: openpyxl (pip install openpyxl).
"""
from __future__ import annotations

import argparse
import csv
import os
import re
import shutil
import sys
from collections import defaultdict
from datetime import datetime

# cp949/cp1252-safe stdout for the Windows users running this.
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

HERE = os.path.dirname(os.path.abspath(__file__))
REGO_DIR = os.path.join(HERE, "out", "rego")

DEFAULT_XLSX     = os.path.join(REGO_DIR, "AustraliaFormat1_2025.xlsx")
DEFAULT_EXISTING = os.path.join(REGO_DIR, "vehicle_tyre_fitment.csv")
DEFAULT_OUTPUT   = os.path.join(REGO_DIR, "vehicle_tyre_fitment.csv")

CARRYOVER_YEAR = 2026       # borrow 2026 rows from the existing CSV


# ─── Column resolver ─────────────────────────────────────────────────────
# The xlsx header row carries human-readable column labels; the exact
# spelling varies between vintages of the source dataset, so we probe
# a short list of synonyms for each logical field.
COL_SYNONYMS = {
    "country":      ("country",),
    "manufacturer": ("manufacturer", "make"),
    "model":        ("model",),
    "trim":         ("trim", "sub-model", "submodel"),
    "version":      ("version", "variant", "modification"),
    "year":         ("year", "gen_start_year", "start_year"),
    "size":         ("size", "tyre", "tyre size", "tire", "tire size"),
    "fa":           ("fa", "front axle", "front_axle"),
    "ra":           ("ra", "rear axle", "rear_axle"),
    # Original/OE column - we try several names.  "Original" is what
    # the user's dataset labels it; wheel-size.com uses "OE".
    "original":     ("original", "oe", "o.e.", "original equipment"),
}


def _resolve_cols(header_row):
    """Return a {logical_name: 0-based_index} map from the xlsx header
    row.  Keys that aren't found map to -1 so callers can safely
    conditional-index them."""
    want = {k: -1 for k in COL_SYNONYMS}
    lowered = [str(h or "").strip().lower() for h in header_row]
    for logical, synonyms in COL_SYNONYMS.items():
        for syn in synonyms:
            syn_l = syn.lower()
            # Exact match first
            for i, h in enumerate(lowered):
                if h == syn_l:
                    want[logical] = i
                    break
            if want[logical] >= 0:
                break
            # Then substring (helps catch "Front Axle Size" etc.)
            for i, h in enumerate(lowered):
                if syn_l in h:
                    want[logical] = i
                    break
            if want[logical] >= 0:
                break
    return want


# ─── FA / RA to position ─────────────────────────────────────────────────
def _axle_to_position(fa, ra) -> str:
    """FA=1 RA=1 -> both; FA=1 RA=0 -> front; FA=0 RA=1 -> rear;
    anything else -> both (treat as fitment-applies-to-this-vehicle,
    matching the existing CSV's default when position was blank)."""
    def _b(v):
        s = str(v or "").strip().lower()
        if s in ("1", "y", "yes", "true", "t"):
            return True
        try:
            return bool(int(float(s)))
        except Exception:
            return False
    f, r = _b(fa), _b(ra)
    if f and not r:   return "front"
    if r and not f:   return "rear"
    return "both"


# ─── Size normaliser - canonical uppercase form ───────────────────────────
_SIZE_RX = re.compile(
    r"\b((?:LT|P)?\d{2,3}/\d{2}R\d{1,2}(?:\.\d)?C?(?:LT)?"
    r"|\d{2,3}[Xx]\d{1,2}(?:\.\d{1,2})?R\d{1,2}(?:LT)?"
    r"|\d{2,3}R\d{1,2}(?:\.\d)?C?)\b",
    re.IGNORECASE,
)


def _norm_size(raw) -> str:
    s = str(raw or "").strip().upper()
    if not s:
        return ""
    m = _SIZE_RX.search(s)
    return m.group(1).upper() if m else s


# ─── Model-name collapser ────────────────────────────────────────────────
def _collapse_model_name(groups):
    """`groups` is a dict keyed by (make, model, year) with value a list
    of {trim, version, size, position} dicts.  Return a list of
    collapsed rows ready for CSV emission:
        {make, model_resolved, year, size, position}

    Collapse policy (per the user's spec):
        * All (trim, version) share the SAME (size, position) set
                                    -> model
        * All versions within each trim converge
                                    -> model-trim
        * Otherwise                 -> model-trim-version
    Each resolved group then emits one row per (size, position) it
    carries.  Resolution runs INDEPENDENTLY for each (make, model,
    year) so the same model-year can collapse differently in another
    year without affecting this one.
    """
    out_rows = []
    for (make, model, year), items in groups.items():
        # Build size-set per (trim, version)
        sz_by_tv = defaultdict(set)
        for it in items:
            tv = (it["trim"], it["version"])
            sz_by_tv[tv].add((it["size"], it["position"]))

        # Does every (trim, version) share the SAME size set?
        unique_sz_sets = {frozenset(v) for v in sz_by_tv.values()}
        if len(unique_sz_sets) == 1:
            resolved = model
            shared = next(iter(unique_sz_sets))
            for (sz, pos) in shared:
                out_rows.append({"make": make, "model": resolved,
                                  "gen_start_year": year,
                                  "size": sz, "position": pos})
            continue

        # Does every version within each trim converge?
        sz_by_trim = defaultdict(set)
        for (t, v), sz_set in sz_by_tv.items():
            sz_by_trim[t].add(frozenset(sz_set))
        if all(len(s) == 1 for s in sz_by_trim.values()):
            for t, sz_set_wrap in sz_by_trim.items():
                # unpack single-element set-of-frozensets to the inner
                # frozenset of (size, pos) tuples
                sz_set = next(iter(sz_set_wrap))
                resolved = f"{model}-{t}" if str(t).strip() else model
                for (sz, pos) in sz_set:
                    out_rows.append({"make": make, "model": resolved,
                                      "gen_start_year": year,
                                      "size": sz, "position": pos})
            continue

        # Versions differentiate - need full model-trim-version
        for (t, v), sz_set in sz_by_tv.items():
            parts = [model]
            if str(t).strip(): parts.append(str(t).strip())
            if str(v).strip(): parts.append(str(v).strip())
            resolved = "-".join(parts)
            for (sz, pos) in sz_set:
                out_rows.append({"make": make, "model": resolved,
                                  "gen_start_year": year,
                                  "size": sz, "position": pos})
    return out_rows


# ─── xlsx reader ─────────────────────────────────────────────────────────
def read_xlsx(path, country_only="AUSTRALIA", oe_only=True):
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("[fatal] openpyxl not installed.  Install with:\n"
              "    pip install openpyxl", file=sys.stderr)
        sys.exit(3)
    print(f"[xlsx] opening {path}", file=sys.stderr)
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    header = next(rows_iter, None)
    if not header:
        print(f"[fatal] {path} has no header row", file=sys.stderr)
        sys.exit(4)
    cols = _resolve_cols(header)
    for req in ("manufacturer", "model", "year", "size"):
        if cols[req] < 0:
            print(f"[fatal] required column missing: {req}  "
                  f"(header: {list(header)})", file=sys.stderr)
            sys.exit(5)
    print(f"[xlsx] column mapping: {cols}", file=sys.stderr)
    print(f"[xlsx] filters: country={country_only!r}  OE-only={oe_only}",
          file=sys.stderr)
    kept = []
    rows_read = rows_dropped_country = rows_dropped_oe = rows_dropped_size = 0
    for r in rows_iter:
        if not r: continue
        rows_read += 1
        def _g(key):
            i = cols[key]
            return r[i] if 0 <= i < len(r) else None
        country = str(_g("country") or "").strip().upper()
        if country_only and country and country != country_only:
            rows_dropped_country += 1
            continue
        if oe_only and cols["original"] >= 0:
            orig = _g("original")
            # Accept 1 / '1' / 'Y' / 'YES' / True; drop everything else.
            orig_s = str(orig or "").strip().lower()
            if orig_s not in ("1", "y", "yes", "true", "t"):
                try:
                    if int(float(orig_s or "0")) != 1:
                        rows_dropped_oe += 1
                        continue
                except Exception:
                    rows_dropped_oe += 1
                    continue
        make  = str(_g("manufacturer") or "").strip().upper()
        model = str(_g("model") or "").strip().upper()
        trim  = str(_g("trim") or "").strip().upper()
        ver   = str(_g("version") or "").strip().upper()
        year  = _g("year")
        try:
            year_i = int(float(str(year).strip()))
        except Exception:
            continue
        size  = _norm_size(_g("size"))
        if not size:
            rows_dropped_size += 1
            continue
        if not make or not model:
            continue
        pos = _axle_to_position(_g("fa"), _g("ra"))
        kept.append({"make": make, "model": model,
                      "trim": trim, "version": ver,
                      "year": year_i,
                      "size": size, "position": pos})
    print(f"[xlsx] rows_read={rows_read} "
          f"dropped(country={rows_dropped_country}, "
          f"oe={rows_dropped_oe}, "
          f"size={rows_dropped_size}) "
          f"kept={len(kept)}", file=sys.stderr)
    return kept


# ─── Existing CSV reader (for 2026 carryover) ─────────────────────────────
def read_existing_for_year(path, year):
    if not os.path.exists(path):
        print(f"[existing] {path} not found - skipping carryover",
              file=sys.stderr)
        return []
    kept = []
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        rdr = csv.DictReader(f)
        for r in rdr:
            try:
                yr = int(float(str(r.get("gen_start_year") or "").strip()))
            except Exception:
                continue
            if yr != year:
                continue
            kept.append({
                "make":           str(r.get("make") or "").strip().upper(),
                "model":          str(r.get("model") or "").strip().upper(),
                "gen_start_year": yr,
                "size":           _norm_size(r.get("size")),
                "position":       (str(r.get("position") or "both")
                                   .strip().lower() or "both"),
            })
    print(f"[existing] carrying over {len(kept)} rows for year={year}",
          file=sys.stderr)
    return kept


# ─── Main ────────────────────────────────────────────────────────────────
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx",         default=DEFAULT_XLSX,
                    help=f"Input xlsx (default: {DEFAULT_XLSX})")
    ap.add_argument("--existing-csv", default=DEFAULT_EXISTING,
                    help=f"Existing CSV for 2026 carryover "
                         f"(default: {DEFAULT_EXISTING})")
    ap.add_argument("--output",       default=DEFAULT_OUTPUT,
                    help=f"Output CSV path (default: {DEFAULT_OUTPUT})")
    ap.add_argument("--carryover-year", type=int, default=CARRYOVER_YEAR,
                    help=f"Year to carry over from the existing CSV "
                         f"(default: {CARRYOVER_YEAR})")
    ap.add_argument("--no-country-filter", action="store_true",
                    help="Keep rows from every Country (default: "
                         "Australia only)")
    ap.add_argument("--include-non-oe", action="store_true",
                    help="Keep aftermarket/optional fitments too "
                         "(default: OE-only / Original=1)")
    ap.add_argument("--dry-run", action="store_true",
                    help="Compute and print stats but DON'T touch the "
                         "output CSV or create a backup.")
    args = ap.parse_args()

    if not os.path.exists(args.xlsx):
        print(f"[fatal] xlsx not found: {args.xlsx}", file=sys.stderr)
        sys.exit(1)

    # 1. Read xlsx + collapse
    xlsx_rows = read_xlsx(
        args.xlsx,
        country_only="" if args.no_country_filter else "AUSTRALIA",
        oe_only=not args.include_non_oe,
    )
    # Group by (make, model, year) so the collapse policy applies.
    groups = defaultdict(list)
    for r in xlsx_rows:
        groups[(r["make"], r["model"], r["year"])].append(r)
    collapsed = _collapse_model_name(groups)
    print(f"[collapse] {len(xlsx_rows)} raw rows "
          f"-> {len(collapsed)} collapsed rows "
          f"across {len(groups)} (make,model,year) groups",
          file=sys.stderr)

    # 2. Carry over 2026 rows from existing CSV
    carryover = read_existing_for_year(args.existing_csv, args.carryover_year)

    # 3. Merge + dedup on (make, model, year, size, position)
    out_rows_by_key = {}
    for r in collapsed + carryover:
        k = (r["make"], r["model"], r["gen_start_year"],
             r["size"], r["position"])
        out_rows_by_key.setdefault(k, r)   # first write wins; both are canonical
    out_rows = sorted(out_rows_by_key.values(),
                       key=lambda r: (r["make"], r["model"],
                                      r["gen_start_year"],
                                      r["size"], r["position"]))
    print(f"[merge] final: {len(out_rows)} unique rows "
          f"(carryover {args.carryover_year}: {len(carryover)}; "
          f"xlsx: {len(collapsed)})", file=sys.stderr)

    if args.dry_run:
        print("[dry-run] not writing anything.  First 10 output rows:",
              file=sys.stderr)
        for r in out_rows[:10]:
            print(f"    {r}", file=sys.stderr)
        return

    # 4. Backup existing CSV (timestamped)
    if os.path.exists(args.output):
        stamp = datetime.now().strftime("%Y%m%d_%H%M")
        bak = f"{args.output}.bak_{stamp}"
        shutil.copy2(args.output, bak)
        print(f"[backup] {args.output} -> {bak}", file=sys.stderr)

    # 5. Write consolidated CSV
    out_dir = os.path.dirname(os.path.abspath(args.output))
    if out_dir: os.makedirs(out_dir, exist_ok=True)
    with open(args.output, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["make", "model", "gen_start_year", "size", "position"])
        for r in out_rows:
            w.writerow([r["make"], r["model"], r["gen_start_year"],
                        r["size"], r["position"]])
    print(f"[done] wrote {len(out_rows)} rows -> {args.output}")


if __name__ == "__main__":
    main()
